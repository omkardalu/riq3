"""Builders for real PDF/DOCX fixtures used across tests."""
import io

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas


def make_pdf(lines, pages=1) -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    for _ in range(pages):
        y = 800
        for line in lines:
            c.drawString(50, y, line)
            y -= 18
        c.showPage()
    c.save()
    return buf.getvalue()


def make_blank_pdf() -> bytes:
    """A PDF with no text layer (what a scanned resume looks like to a text extractor)."""
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    c.rect(50, 50, 200, 200, fill=1)  # graphics only
    c.showPage()
    c.save()
    return buf.getvalue()


def make_docx(lines, table=None) -> bytes:
    d = Document()
    for line in lines:
        d.add_paragraph(line)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for i, row in enumerate(table):
            for j, cell in enumerate(row):
                t.cell(i, j).text = cell
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


SAMPLE_LINES = [
    "Rahul Kumar", "rahul@example.com  +91 98765 43210", "",
    "Skills", "Python, JavaScript, SQL, Git, REST API",
    "Experience", "Frontend Developer Intern - 6 months. Built responsive web applications.",
    "Education", "B.Tech in Computer Science, XYZ Engineering College, 2027",
    "Projects", "AI Resume Analyzer; Student Management System",
]
